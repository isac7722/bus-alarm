from __future__ import annotations

import asyncio
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import structlog
from openpyxl import load_workbook
from sqlalchemy import delete, insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.database import async_session_factory
from app.models.catalog import Route, RouteStop, Station

EXPECTED_HEADERS = ("ROUTE_ID", "노선명", "순번", "NODE_ID", "ARS_ID", "정류소명", "X좌표", "Y좌표")
INSERT_BATCH_SIZE = 5_000
logger = structlog.get_logger(__name__)


class CatalogImportError(ValueError):
    """Raised when the workbook cannot be safely normalized."""


@dataclass(frozen=True, slots=True)
class CatalogData:
    """Validated normalized rows ready for bulk persistence."""

    stations: tuple[dict[str, Any], ...]
    routes: tuple[dict[str, Any], ...]
    route_stops: tuple[dict[str, Any], ...]
    skipped_non_seoul_rows: int


def read_catalog(path: Path) -> CatalogData:
    """Read and validate the source workbook without modifying it."""
    if not path.is_file():
        raise CatalogImportError(f"Catalog file does not exist: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Data" not in workbook.sheetnames:
        workbook.close()
        raise CatalogImportError("Catalog workbook must contain a 'Data' sheet")
    worksheet = workbook["Data"]
    rows = worksheet.iter_rows(values_only=True)
    headers = tuple(next(rows, ()))
    if headers != EXPECTED_HEADERS:
        workbook.close()
        raise CatalogImportError(f"Unexpected headers: {headers!r}")

    stations: dict[str, dict[str, Any]] = {}
    routes: dict[str, dict[str, Any]] = {}
    route_stops: dict[tuple[str, int], dict[str, Any]] = {}
    skipped_non_seoul_rows = 0

    for row_number, row in enumerate(rows, start=2):
        if len(row) != len(EXPECTED_HEADERS) or any(value is None for value in row):
            workbook.close()
            raise CatalogImportError(f"Row {row_number} has missing or malformed values")
        route_id, route_name, sequence, node_id, ars_id, station_name, longitude, latitude = row
        normalized_node_id = str(node_id).strip()
        if not normalized_node_id.startswith("1"):
            skipped_non_seoul_rows += 1
            continue
        normalized_station_id = str(ars_id).strip().split(".")[0].zfill(5)
        normalized_route_id = str(route_id).strip().split(".")[0]
        if len(normalized_station_id) != 5 or not normalized_station_id.isdigit():
            workbook.close()
            raise CatalogImportError(f"Row {row_number} has an invalid ARS_ID")
        if len(normalized_node_id) != 9 or not normalized_node_id.isdigit():
            workbook.close()
            raise CatalogImportError(f"Row {row_number} has an invalid NODE_ID")
        try:
            normalized_sequence = int(sequence)
            normalized_longitude = float(longitude)
            normalized_latitude = float(latitude)
        except (TypeError, ValueError) as error:
            workbook.close()
            raise CatalogImportError(f"Row {row_number} has invalid numeric values") from error
        if normalized_sequence < 1 or not 123 <= normalized_longitude <= 133 or not 32 <= normalized_latitude <= 40:
            workbook.close()
            raise CatalogImportError(f"Row {row_number} contains out-of-range data")

        station = {
            "station_id": normalized_station_id,
            "node_id": normalized_node_id,
            "name": str(station_name).strip(),
            "longitude": normalized_longitude,
            "latitude": normalized_latitude,
        }
        existing_station = stations.get(normalized_station_id)
        if existing_station is not None and existing_station != station:
            workbook.close()
            raise CatalogImportError(f"Conflicting station metadata for ARS_ID {normalized_station_id}")
        stations[normalized_station_id] = station

        route = {"route_id": normalized_route_id, "name": str(route_name).strip()}
        existing_route = routes.get(normalized_route_id)
        if existing_route is not None and existing_route != route:
            workbook.close()
            raise CatalogImportError(f"Conflicting route metadata for route {normalized_route_id}")
        routes[normalized_route_id] = route

        route_stop_key = (normalized_route_id, normalized_sequence)
        route_stop = {
            "route_id": normalized_route_id,
            "sequence": normalized_sequence,
            "station_id": normalized_station_id,
        }
        if route_stop_key in route_stops:
            workbook.close()
            raise CatalogImportError(f"Duplicate route sequence at row {row_number}: {route_stop_key}")
        route_stops[route_stop_key] = route_stop

    workbook.close()
    if not stations or not routes or not route_stops:
        raise CatalogImportError("Catalog contained no Seoul station rows")
    return CatalogData(
        stations=tuple(stations.values()),
        routes=tuple(routes.values()),
        route_stops=tuple(route_stops.values()),
        skipped_non_seoul_rows=skipped_non_seoul_rows,
    )


def batched(rows: Sequence[dict[str, Any]], size: int = INSERT_BATCH_SIZE) -> Iterable[Sequence[dict[str, Any]]]:
    """Yield bounded bulk-insert batches."""
    for start in range(0, len(rows), size):
        yield rows[start : start + size]


async def persist_catalog(session: AsyncSession, catalog: CatalogData) -> None:
    """Atomically replace the catalog with validated workbook content."""
    async with session.begin():
        await session.execute(delete(RouteStop))
        await session.execute(delete(Route))
        await session.execute(delete(Station))
        for batch in batched(catalog.stations):
            await session.execute(insert(Station), batch)
        for batch in batched(catalog.routes):
            await session.execute(insert(Route), batch)
        for batch in batched(catalog.route_stops):
            await session.execute(insert(RouteStop), batch)


async def run_import(path: Path) -> CatalogData:
    """Load a workbook and replace the database catalog."""
    catalog = await asyncio.to_thread(read_catalog, path)
    async with async_session_factory() as session:
        await persist_catalog(session, catalog)
    logger.info(
        "station_catalog_imported",
        stations=len(catalog.stations),
        routes=len(catalog.routes),
        route_stops=len(catalog.route_stops),
        skipped_non_seoul_rows=catalog.skipped_non_seoul_rows,
    )
    return catalog


def main() -> None:
    """CLI entry point using the configured default workbook path."""
    asyncio.run(run_import(Path(get_settings().station_catalog_path)))


if __name__ == "__main__":
    main()
